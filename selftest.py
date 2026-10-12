# -*- coding: utf-8 -*-
"""自动自检（不需要界面）：每次改完先跑这个。

    python selftest.py            # 全部检查
    python selftest.py --quick    # 跳过慢的（工作簿导出、全量候选池扫描）

界面相关的检查在 `selftest_ui.py`（需要能开窗口的环境）。
两套合起来就是 `TESTING.md` 里的「阶段 1 / 阶段 2」。
"""
import argparse
import os
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import engine  # noqa: E402
import run  # noqa: E402

if getattr(sys, "stdout", None):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

TOL = 1e-9
RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), detail))
    print("  %s %s%s" % ("PASS" if ok else "FAIL", name,
                         ("　— " + detail) if detail else ""))
    return ok


def approx(a, b, tol=1e-6):
    return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))


def inputs_for(country="SOV", role="air_far", tech_year=1944, current_year=1944, **kw):
    i = dict(run.DEFAULT_INPUTS)
    i.update({"country": country, "role": role, "tech_year": tech_year,
              "current_year": current_year, "roles": [role],
              "enemy": {"country": "GER" if country != "GER" else "SOV",
                        "rule": "previous_year", "override_key": None, "by_country": {}}})
    i.update(kw)
    return i


# ---------------------------------------------------------------- 1. 锚点
def t_anchor():
    print("[1] 验收锚点（Yak-9U@1944 对空 900）")
    code = run.anchor_check()
    return check("锚点校验", code == 0, "run.anchor_check() 返回 %d" % code)


# ---------------------------------------------------------------- 2. 候选池
def t_pool_window(data, quick=False):
    print("[2] 候选池「最近 3 个年式」")
    ok = True
    combos = [("SOV", "air_far", 1944, 40), ("SOV", "air_far", 1942, 32)]
    if not quick:
        combos += [("GER", "air_far", 1948, None), ("ITA", "recon", 1944, None)]
    for country, role, ty, old_n in combos:
        i = inputs_for(country, role, ty)
        pool = run.candidate_pool(i, data)
        if not pool:
            ok &= check("%s %s %d 候选池非空" % (country, role, ty), False, "空池")
            continue
        years = [a.get("year") or 0 for a in pool]
        newest = max(years)
        ok &= check("%s %s %d 只含最近 3 个年式" % (country, role, ty),
                    min(years) >= newest - (run.POOL_WINDOW - 1),
                    "年份 %s" % sorted(set(years)))
        if old_n:                            # 老规则下的架数，用来确认"确实收窄了"
            ok &= check("%s %s %d 比旧规则少" % (country, role, ty), len(pool) < old_n,
                        "新 %d < 旧 %d" % (len(pool), old_n))
    return ok


# ------------------------------------------------------- 3. 资源加权分恒等式
def t_weighted(data):
    print("[3] 资源加权分 = 折算后总分数 ÷ 资源开销")
    ok = True
    for alloc in ("split", "balanced"):
        i = inputs_for("SOV", "air_far", 1944, 1944,
                       manufacturer_mode="traits", modifier_source="mio", alloc_mode=alloc,
                       unlocked_techs=["advanced_machine_tools", "computing_machine",
                                       "worker_training", "assembly_line_production",
                                       "Ltaircraft_industry2", "Ltaircraft_industry3",
                                       "chemical_industry_iii", "improved_centimetric_radar",
                                       "advanced_centimetric_radar"],
                       unlocked_flags=["unlock_plane_anti_air_rocket_upgrade"],
                       mio={"organization": "SOV_lavochkin_design_bureau_organization",
                            "level": 14, "manual_traits": None,
                            "policy": "mio_policy_air_wind_tunnel_pioneers"})
        rows, _m, _e, _ev, mfr = run.build_rows(i, data)
        res = [r for r in rows if "9U" in ((r["airframe"] or {}).get("name") or "")]
        if not res:
            ok &= check("%s 找到 Yak-9U 行" % alloc, False)
            continue
        r = res[0]
        need = (r.get("production") or {}).get("production_resource_need_factor", 0.0)
        cost = (r["airframe"].get("air_production") or 0) * (1.0 + need)
        ok &= check("%s 恒等式（%s）" % (alloc, r["airframe"]["name"]),
                    cost and approx(r["weighted"], r["score"] / cost),
                    "weighted %.4f = score %.4f ÷ %.4f" % (r["weighted"], r["score"], cost))
        ok &= check("%s 分数含生产折算" % alloc,
                    approx(r["score"], r["score_raw"] * r["production_gain"]),
                    "gain %.6f" % r["production_gain"])
    return ok


# ---------------------------------------------------------------- 4. 生产折算
def t_production(data):
    print("[4] 生产折算公式")
    gain, parts = engine.production_gain({"production_capacity_factor": 0.05}, 100, 200)
    ok = check("产量 +5% 在 100/200 上 = ×1.016667", approx(gain, 1.0166666, 1e-6),
               "gain %.6f" % gain)
    gain2, _ = engine.production_gain({"production_cost_factor": -0.03}, 100, 200)
    ok &= check("生产产出花费 −3% = ×1.030928", approx(gain2, 1.0309278, 1e-6),
                "gain %.6f" % gain2)
    gain3, _ = engine.production_gain({"production_capacity_factor": 0.05,
                                       "production_cost_factor": -0.03}, 100, 200)
    ok &= check("两项可乘（×1.0167 × ×1.0309）", approx(gain3, gain * gain2, 1e-9),
                "gain %.6f" % gain3)
    return ok


# ------------------------------------------------------------ 5. 对手 AI 判定
def t_enemy(data):
    print("[5] 对手（AI）设计与经验")
    _e, values = engine.pick_enemy_values(data["airframes_all"],
                                         {"country": "GER", "rule": "previous_year"},
                                         "air_far", 1944)
    ok = check("德国 1944 对空 = 机动 86.112 / 速度 708.72375",
               values and approx(values.get("agility", 0), 86.112, 1e-6)
               and approx(values.get("speed", 0), 708.72375, 1e-6),
               "%.5f / %.5f" % ((values or {}).get("agility", 0), (values or {}).get("speed", 0)))
    _e2, cv = engine.pick_enemy_values(data["airframes_all"],
                                       {"country": "GER", "rule": "previous_year"},
                                       "air_cv", 1944)
    ok &= check("舰载对手不计经验（数值低于陆基口径）",
                cv and cv.get("agility", 0) < values.get("agility", 0),
                "舰载 %.2f < 陆基 %.2f" % (cv.get("agility", 0), values.get("agility", 0)))
    return ok


# ------------------------------------------------------------ 6. 命名与版本
def t_range_coverage(data):
    """航程一维覆盖模型（§50）：近边距离 = 最远目标距离 − 目标距离差。"""
    print("[9] 航程覆盖模型（最远目标距离 + 目标距离差）")
    spec = engine.RangeSpec(900, 300)            # 近边 = 600
    cases = [(450, 0.0), (600, 0.0), (750, 0.5), (900, 1.0), (1200, 1.0)]
    ok = True
    for rng, want in cases:
        got = engine.range_coverage(rng, spec)
        ok &= check("航程 %d @ d₁=900/Δ=300 → %.2f" % (rng, want), approx(got, want, 1e-9),
                    "%.4f" % got)
    # Δ = 0 → 区域缩成"最远目标距离上的一个点" → 二段式（航程 ≥ d₁ 才 1）
    point = engine.RangeSpec(900, 0)
    ok &= check("Δ=0 = 点目标（二段式）",
                approx(engine.range_coverage(899, point), 0.0)
                and approx(engine.range_coverage(900, point), 1.0)
                and approx(engine.range_coverage(1200, point), 1.0),
                "899→%.3f，900→%.3f，1200→%.3f" % (engine.range_coverage(899, point),
                                                   engine.range_coverage(900, point),
                                                   engine.range_coverage(1200, point)))
    # Δ = d₁（近边在机场脚下）才是旧口径；Δ > d₁ 也夹回旧口径
    for band, label in ((900, "Δ=d₁"), (1200, "Δ>d₁")):
        old = engine.RangeSpec(900, band)
        ok &= check("%s 退化成旧口径 min(R,d₁)/d₁" % label,
                    approx(engine.range_coverage(450, old), 0.5)
                    and approx(engine.range_coverage(700, old), 700 / 900),
                    "450→%.3f，700→%.3f" % (engine.range_coverage(450, old),
                                            engine.range_coverage(700, old)))
    # 端到端：改「目标距离差」必须真的改变分数（说明它一路传到了引擎）。
    # 基线用 Δ = d₁ = 900（近边在机场脚下 = 旧口径），对照新默认 Δ=300。
    base = inputs_for("SOV", "air_far", 1944, 1944)
    rows_a, *_ = run.build_rows(dict(base, range_band=900), data)
    rows_b, *_ = run.build_rows(dict(base, range_band=300), data)
    fa = {r["airframe"]["key"]: r["values"]["range_factor"]
          for r in rows_a if r["role"] == "air_far"}
    fb = {r["airframe"]["key"]: r["values"]["range_factor"]
          for r in rows_b if r["role"] == "air_far"}
    sa = {r["airframe"]["key"]: r["score"] for r in rows_a if r["role"] == "air_far"}
    sb = {r["airframe"]["key"]: r["score"] for r in rows_b if r["role"] == "air_far"}
    lowered = [k for k in fa if fb.get(k, 1.0) < fa[k] - 1e-9]
    changed = [k for k in sa if abs(sa[k] - sb.get(k, 0.0)) > 1e-9]
    top_key = next(r for r in rows_b if r["role"] == "air_far")["airframe"]["key"]
    ok &= check("改 Δ 会改变分数", bool(changed),
                "%d/%d 架分数不同；首名 %.4f（Δ=900 时 %.4f）"
                % (len(changed), len(sa), sb[top_key], sa.get(top_key, 0.0)))
    ok &= check("Δ=300 会压低短航程飞机的覆盖率", bool(lowered),
                "%d 架覆盖率被压低（其中 %d 架为 0）"
                % (len(lowered), sum(1 for k in lowered if fb[k] <= 0)))
    # 起步（2026-10-12 修）：裸机航程 < 近边、但加航程改装后能到的机型，不该被判 0 分。
    # 1941/科技 1942 @900/300：修前 LaGG-1（588）、MiG-1（571）都是 0 分。
    rows_c, *_ = run.build_rows(inputs_for("SOV", "air_far", 1942, 1941), data)
    rc = [r for r in rows_c if r["role"] == "air_far"]
    zero_c = [r for r in rc if r["values"]["range_factor"] <= 0]
    ok &= check("裸机够不到、加改装能到的机型不再被判 0（1941/科技1942）", not zero_c,
                "0 分：%s" % [r["airframe"]["name"] for r in zero_c])
    # 远距离口径：够得到的要有分、够不到的就得是 0（1946 @ 最远 1200 / 差值 300）
    rows_d, *_ = run.build_rows(inputs_for("SOV", "air_far", 1946, 1946,
                                           range_cap=1200.0, range_band=300.0), data)
    rd = [r for r in rows_d if r["role"] == "air_far"]
    top_d = rd[0]
    blocked = sum(1 for r in rd if r["values"]["range_factor"] <= 0)
    ok &= check("1946 @ 最远 1200/差值 300：够得到的有分、够不到的 0 分",
                top_d["score"] > 0 and top_d["values"]["range_factor"] > 0 and blocked > 0,
                "首名 %s 分 %.3f 覆盖率 %.3f；0 分 %d/%d 架"
                % (top_d["airframe"]["name"], top_d["score"],
                   top_d["values"]["range_factor"], blocked, len(rd)))
    # 多步起步（2026-10-12 二次修）：裸机分 0、且**单步过不了线**的机型也不能被判 0 ——
    # 航程改装里有"平添"部分（增程 +30/级、副油箱 +120），常常要连点两三级才够到近边。
    # 场景：无 MIO 的苏联 1939 @900/300（近边 600）：修前 7 架里 5 架是 0。
    base6 = inputs_for("SOV", "air_far", 1939, 1939, manufacturer_mode="manual",
                       manufacturer_modifiers={}, modifier_source="manual")
    rows_e, *_ = run.build_rows(base6, data)
    re_ = [r for r in rows_e if r["role"] == "air_far"]
    top_e = re_[0]
    ok &= check("多步起步：连点航程改装能过线的机型要有正覆盖率",
                top_e["score"] > 0 and top_e["values"]["range_factor"] > 0
                and ("燃料容量" in top_e["design"] or "副油箱" in top_e["design"]),
                "首名 %s 分 %.3f 覆盖率 %.3f 配方 %s"
                % (top_e["airframe"]["name"], top_e["score"],
                   top_e["values"]["range_factor"], top_e["design"]))
    # 用户实测复现：苏联 1944 / 科技 1946 / 最远 1200（近边 900）/ 航程 +4%（拉沃契金档）
    # —— MiG-9 修前 0 分（贪心一步不加），修后应拿到正覆盖率（942 km → 0.14）。
    mig = inputs_for("SOV", "air_far", 1946, 1944, range_cap=1200.0, range_band=300.0,
                     manufacturer_mode="manual", manufacturer_modifiers={"air_range": 0.04},
                     modifier_source="manual")
    rows_f, *_ = run.build_rows(mig, data)
    rf = [r for r in rows_f if "Mikoyan-Gurevich 9" in r["airframe"]["name"]]
    ok &= check("MiG-9 @1946/最远1200 不再是 0 分",
                bool(rf) and rf[0]["values"]["range_factor"] > 0,
                ("覆盖率 %.3f 航程 %.0f" % (rf[0]["values"]["range_factor"], rf[0]["values"]["range"])
                 if rf else "候选里没找到 MiG-9"))
    return ok


def t_naming(data):
    print("[6] 版本标识与导出文件名")
    info = run.build_info()
    label = run.edition_label()
    ok = check("data/build_info.json 可读且标题非空", bool(label) and "未标注" not in label,
               "edition_label() = %r" % label)
    ok &= check("版本号与 exe 名一致（%s）" % info.get("mod_version"),
                bool(info.get("mod_version")), "workshop %s" % info.get("workshop_id"))
    i = inputs_for("SOV", "air_far", 1944, 1944)
    p = run.output_path(i)
    base = os.path.basename(p)
    ok &= check("导出文件名带 国家/年份/用途", all(s in base for s in ("SOV", "1944", "对空")),
                base)
    ok &= check("导出目录 = <app>/output", os.path.dirname(p) == run.OUTPUT_DIR,
                os.path.dirname(p))
    i2 = inputs_for("GER", "air_cv", 1943, 1942)
    ok &= check("多用途/舰载命名正常", "对空舰载" in os.path.basename(run.output_path(i2)),
                os.path.basename(run.output_path(i2)))
    return ok


# ------------------------------------------------------------ 7. 离线表分片
def t_table_files(data):
    print("[7] 离线表分片")
    import glob
    import json
    files = sorted(glob.glob(os.path.join(run.PLAN_TABLE_DIR, "*.json")))
    total = 0
    for p in files:
        with open(p, encoding="utf-8") as fh:
            total += len(json.load(fh))
    ok = check("12 片都在", len(files) == 12, "%d 片" % len(files))
    ok &= check("条目数 ≥ 6800", total >= 6800, "%d 条" % total)
    return ok


# ------------------------------------------------------------ 8. 导出工作簿
def t_workbook(data):
    print("[8] 导出工作簿（写临时文件）")
    i = inputs_for("SOV", "air_far", 1944, 1944)
    rows, mods, enemy, ev, mfr = run.build_rows(i, data)
    path = os.path.join(tempfile.gettempdir(), "_selftest_out.xlsx")
    if os.path.exists(path):
        os.remove(path)
    try:
        run.write_workbook(path, i, rows, mods, enemy, ev, mfr)
    except Exception as exc:                 # noqa: BLE001
        return check("write_workbook 不报错", False, repr(exc))
    ok = check("写出文件", os.path.exists(path) and os.path.getsize(path) > 5000,
               "%d 字节" % (os.path.getsize(path) if os.path.exists(path) else 0))
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ok &= check("两张表（选型 + 参数与口径）", len(wb.sheetnames) >= 2,
                    ",".join(wb.sheetnames))
    except Exception as exc:                 # noqa: BLE001
        ok &= check("openpyxl 能打开", False, repr(exc))
    try:
        os.remove(path)
    except OSError:
        pass
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--quick", action="store_true", help="跳过慢检查")
    args = ap.parse_args()
    t0 = time.time()
    data = engine.load_data()
    print("数据指纹：%s　候选池窗口：%d 年式" % (engine.data_fingerprint(), run.POOL_WINDOW))
    t_anchor()
    t_pool_window(data, args.quick)
    t_weighted(data)
    t_production(data)
    t_enemy(data)
    t_range_coverage(data)
    t_naming(data)
    t_table_files(data)
    if not args.quick:
        t_workbook(data)
    bad = [n for n, ok, _d in RESULTS if not ok]
    print("\n合计 %d 项，失败 %d 项，用时 %.1fs" % (len(RESULTS), len(bad), time.time() - t0))
    if bad:
        print("失败：" + "；".join(bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
